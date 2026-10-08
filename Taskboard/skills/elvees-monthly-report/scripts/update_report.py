#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Edit the MAIN FPI report in place: 'Отчетность в ФПИ. Шаблон отчёта ФПИ.docx'.

Copy the previous month's template into the new output dir first, then run.
Preserves formatting by editing runs; keeps the example row untouched.

Rules:
- Title date -> current PERIOD.
- Each task row: col2 = ТЗ-пункт + название + 2 SHORT sentences of work;
  col3 = ARCHIVE NAME (must equal the real .zip on disk); col4/col5 = % from ОРСПО.
- Description: ровно «одно предложение что это + одно что сделано» (Ilya: two simple sentences).

Inspect the template's actual row indices first (they vary); this expects:
rows 0-1 header/numbering, row 2 example, rows 3.. = task rows.
"""
import docx

PATH = "/home/auser/src/elvees/other/MR_report/output-2026-09/Отчетность в ФПИ. Шаблон отчёта ФПИ.docx"
OLD_PERIOD = ("01.08.2026", "31.08.2026")
NEW_PERIOD = ("01.09.2026", "30.09.2026")

# row index in table -> fields
ROWS = {
    3: {  # Восход
        "p0": "2.2.3. Разработка программного обеспечения макета Платформы (п. 4.1.7).",
        "p2": "Развитие архитектуры управления вычислительными ядрами и радиоприложениями.",
        "p3": "\nВведены долгоживущий CoreManager и RadioManager, компилирующий граф в поединичные "
              "инструкции. Добавлены единый бинарник bundledRM, развёртывание по ConnectionManifest "
              "и централизованный сбор логов runtime.",
        "zip": "RadioManager_architecture.zip",
        "pct": "86%", "done": "100%",
    },
    4: {  # Дебют (single merged paragraph p2 style in Aug template)
        "p0": "ОРСПО, строка 246.",
        "p2": "Сравнение быстродействия mcom03_ipc и dds middleware\n\n"
              "Доработаны два межъядерных транспорта middleware: канал mcom03_ipc (MCOM03 mux) и канал "
              "Fast DDS/XRCE. Добавлен TCP-транспорт, бинарные логи DSP переведены с DDS-канала на IPC.",
        "zip": "IPC_vs_DDS_middleware.zip",
        "pct": "14%", "done": "100%",
    },
}


def set_para(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def main():
    d = docx.Document(PATH)
    for r in d.paragraphs[0].runs:
        if OLD_PERIOD[0] in r.text:
            r.text = r.text.replace(OLD_PERIOD[0], NEW_PERIOD[0]).replace(OLD_PERIOD[1], NEW_PERIOD[1])
    t = d.tables[0]
    for ri, f in ROWS.items():
        c2 = t.rows[ri].cells[2]
        paras = c2.paragraphs
        if "p0" in f and len(paras) > 0: set_para(paras[0], f["p0"])
        if "p2" in f and len(paras) > 2: set_para(paras[2], f["p2"])
        if "p3" in f and len(paras) > 3: set_para(paras[3], f["p3"])
        c3 = t.rows[ri].cells[3]
        if len(c3.paragraphs) > 1:
            set_para(c3.paragraphs[1], f["zip"])
        # percentages (col4/col5) — set only the first run to keep formatting
        for col, key in ((4, "pct"), (5, "done")):
            cell = t.rows[ri].cells[col]
            if cell.paragraphs and cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].text = f[key]
                for r in cell.paragraphs[0].runs[1:]:
                    r.text = ""
    d.save(PATH)
    print("updated", PATH)


if __name__ == "__main__":
    main()
