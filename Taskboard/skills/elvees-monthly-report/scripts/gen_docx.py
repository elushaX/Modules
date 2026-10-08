#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the per-task change-description DOCX (Изменения_<task>_<месяц>.docx).

Run after gen_packages.py (reads /tmp/fpi_meta.json):
    /tmp/fpi_venv/bin/python gen_docx.py

The MAIN FPI report is NOT generated here — it is the in-place edited
'Отчетность в ФПИ. Шаблон отчёта ФПИ.docx' (see update_report.py).
"""
import json, subprocess, os
from docx import Document
from docx.shared import Pt, Cm

REPO = "/home/auser/src/elvees/er_platform/runtime"
MONTH_LABEL = "сентябрь_2026"          # edit per month
MONTH_HUMAN = "сентябрь 2026"
PERIOD = "01.09.2026–30.09.2026"
# task folder -> (snippet commit short sha, max diff lines in docx)
SNIPPETS = {"RadioManager_architecture": ("df835ff6", 34)}

meta = json.load(open("/tmp/fpi_meta.json"))


def diff_snippet(h, n):
    out = subprocess.run(["git", "-C", REPO, "diff", "--no-ext-diff", "--full-index",
                          f"{h}^", h, "--"], capture_output=True, text=True).stdout.splitlines()
    return out[:n]


def build(pkgdir):
    info = meta[pkgdir]; cfg = info["cfg"]; metas = info["metas"]; base = info["base"]
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(2); s.right_margin = Cm(1.5)

    def P(t, bold=False, size=11, after=6):
        p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(after); return p

    P(cfg["readme_title"], bold=True, size=13)
    P(f"Шурупов Илья Викторович, системный программист. Проект «Восход». Период: {PERIOD}.", size=10)
    P(cfg["intro"])
    P(cfg["platforms"])
    for m in metas:
        P(f"{m['short']} — {m['desc']}", bold=True, size=11, after=2)
        P(f"Автор: {m['an']}", size=9, after=0)
        P(f"Дата автора: {m['ad']}", size=9, after=0)
        P(f"Дата коммита: {m['cd']}", size=9, after=0)
        P(m["subj"], size=9, after=0)
        if m["statline"]:
            P(m["statline"], size=9, after=0)
        P(f"Полный diff: {m['difffile']}", size=9, after=8)
    if pkgdir in SNIPPETS:
        sh, n = SNIPPETS[pkgdir]
        full = next(x["full"] for x in metas if x["short"] == sh)
        lines = diff_snippet(full, n)
        if lines:
            hint = lines[0].split(" b/")[-1] if " b/" in lines[0] else ""
            P(f"Фрагмент diff: {sh} / {hint}", bold=True, size=10, after=2)
            for ln in lines:
                p = doc.add_paragraph(); r = p.add_run(ln)
                r.font.name = "Consolas"; r.font.size = Pt(8)
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    P("Границы подтверждения", bold=True, size=11, after=2)
    P("Представлены исходные изменения и, где применимо, код проверок, а не журналы успешных "
      "прогонов. Включены полные неизменённые diff выбранных коммитов относительно их родителей; "
      "SHA-256 и patch-id приведены в manifest.json. Даты всех выбранных изменений в " + MONTH_HUMAN +
      ", автор — ishurupov@elvees.com. Это тематическая выборка за месяц, а не полный совокупный "
      "diff всех веток; применять весь набор подряд к текущему HEAD не следует. Аппаратная "
      "валидация и показатели производительности в данном пакете не заявляются.", size=10)
    name = f"Изменения_{pkgdir}_{MONTH_LABEL}.docx"
    doc.save(os.path.join(base, name))
    return name


def main():
    for pkgdir in meta:
        if pkgdir.startswith("_"):
            continue
        print("change-doc:", build(pkgdir))


if __name__ == "__main__":
    main()
